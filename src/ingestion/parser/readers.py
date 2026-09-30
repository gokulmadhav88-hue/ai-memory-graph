"""One reader per file type (rule 7). Every reader returns a RawDocument:
prose as a list of pages, plus any tables found (rule 8), plus any figures found
(Tier 3), plus which pages (if any) came from OCR instead of a real text layer.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

from . import config
from .figures import FigureRef, ocr_pages
from .tables import Table, extract_captions, is_caption
from .validator import InvalidFile, ParserError

log = logging.getLogger(__name__)


@dataclass
class RawDocument:
    pages: list[str] = field(default_factory=list)
    tables: list[Table] = field(default_factory=list)
    figures: list[FigureRef] = field(default_factory=list)
    ocr_pages: set[int] = field(default_factory=set)   # 1-indexed page numbers read via OCR

    def is_empty(self) -> bool:
        return not any(p.strip() for p in self.pages) and not self.tables


_HEADING_MAX_WORDS = 8


def _looks_like_heading(line: str) -> str | None:
    line = line.strip()
    if not line or "\n" in line:
        return None
    if line.startswith("#"):
        return line.lstrip("#").strip() or None
    if len(line.split()) <= _HEADING_MAX_WORDS and not line.endswith((".", "!", "?", ":", ";", ",")):
        return line
    return None


# ---------- .txt and .md ----------

def _read_txt(path: Path) -> RawDocument:
    return RawDocument(pages=[path.read_text(encoding="utf-8-sig", errors="replace")])


_MD_ROW = re.compile(r"^\s*\|.*\|\s*$")
_MD_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _split_md_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def _extract_md_tables(text: str) -> tuple[str, list[Table]]:
    lines = text.split("\n")
    out: list[str] = []
    tables: list[Table] = []
    current_section: str | None = None
    i = 0
    while i < len(lines):
        heading = _looks_like_heading(lines[i]) if lines[i].strip().startswith("#") else None
        if heading:
            current_section = heading

        if _MD_ROW.match(lines[i]) and i + 1 < len(lines) and _MD_SEP.match(lines[i + 1]):
            rows = [_split_md_row(lines[i])]
            j = i + 2
            while j < len(lines) and _MD_ROW.match(lines[j]):
                rows.append(_split_md_row(lines[j]))
                j += 1
            caption = None
            k = len(out) - 1
            while k >= 0 and not out[k].strip():
                k -= 1
            if k >= 0 and is_caption(out[k]):
                caption = is_caption(out[k])
                out = out[:k]
            tables.append(Table(rows=rows, caption=caption, section=current_section))
            out.append("")
            i = j
        else:
            out.append(lines[i])
            i += 1
    return "\n".join(out), tables


def _read_markdown(path: Path) -> RawDocument:
    text, tables = _extract_md_tables(path.read_text(encoding="utf-8-sig", errors="replace"))
    return RawDocument(pages=[text], tables=tables)


# ---------- .pdf ----------

def _pdf_page(page, number: int) -> tuple[str, list[Table]]:
    try:
        accepted = []
        for t in page.find_tables():
            rows = t.extract()
            if len(rows) >= 2 and max(len(r) for r in rows) >= 2:
                accepted.append((t, rows))
        if not accepted:
            return page.extract_text() or "", []

        view = page
        for t, _ in accepted:
            view = view.outside_bbox(t.bbox)
        text, captions = extract_captions(view.extract_text() or "", len(accepted))
        tables = [Table(rows=rows, caption=captions[i], page=number)
                  for i, (_, rows) in enumerate(accepted)]
        return text, tables
    except Exception as e:
        log.warning("page %d: table detection failed (%s), reading it as plain text", number, e)
        return page.extract_text() or "", []


def _pdf_figures(path: Path) -> list[FigureRef]:
    """Embedded raster images, via PyMuPDF (pdfplumber can count images but not
    cleanly hand back their bytes). Isolated in its own try/except: if this fails,
    the document is still readable, it just has no figure descriptions."""
    try:
        import pymupdf
    except Exception as e:
        log.warning("PyMuPDF not available (%s), skipping figure extraction", e)
        return []

    figures: list[FigureRef] = []
    try:
        with pymupdf.open(path) as doc:
            for page_number, page in enumerate(doc, start=1):
                for img in page.get_images(full=True):
                    xref = img[0]
                    try:
                        base = doc.extract_image(xref)
                    except Exception:
                        continue
                    figures.append(FigureRef(
                        image_bytes=base["image"], page=page_number,
                        width=base.get("width", 0), height=base.get("height", 0),
                    ))
    except Exception as e:
        log.warning("%s: figure extraction failed (%s), continuing without figures", path.name, e)
    return figures


def _render_pdf_pages(path: Path) -> list[bytes]:
    """Render each page to a PNG image, for OCR on a scanned/image-only PDF."""
    import pymupdf

    images: list[bytes] = []
    with pymupdf.open(path) as doc:
        zoom = config.OCR_RENDER_DPI / 72
        matrix = pymupdf.Matrix(zoom, zoom)
        for page in doc:
            pix = page.get_pixmap(matrix=matrix)
            images.append(pix.tobytes("png"))
    return images


def _read_pdf(path: Path) -> RawDocument:
    import pdfplumber

    try:
        pdf = pdfplumber.open(path)
    except Exception as e:
        raise InvalidFile(f"{path.name}: cannot open PDF ({e})") from e

    pages: list[str] = []
    tables: list[Table] = []
    bad_pages = 0

    with pdf:
        for number, page in enumerate(pdf.pages, start=1):
            try:
                text, page_tables = _pdf_page(page, number)
            except Exception as e:
                log.warning("%s: page %d could not be read (%s), skipping", path.name, number, e)
                bad_pages += 1
                text, page_tables = "", []
            pages.append(text)
            tables.extend(page_tables)

    figures = _pdf_figures(path)
    doc = RawDocument(pages=pages, tables=tables, figures=figures)

    if doc.is_empty():                                   # rule 4: blank or scanned
        if config.ENABLE_OCR:
            log.info("%s: no extractable text, running OCR", path.name)
            rendered = _render_pdf_pages(path)
            ocr_text = ocr_pages(rendered)
            kept_pages = [t if len(t.strip()) >= config.OCR_MIN_CHARS else "" for t in ocr_text]
            if not any(p.strip() for p in kept_pages):
                log.warning("%s: OCR found no usable text either, skipping", path.name)
                return RawDocument()
            return RawDocument(
                pages=kept_pages,
                ocr_pages={i + 1 for i, p in enumerate(kept_pages) if p.strip()},
            )
        log.warning("%s: no extractable text (blank or scanned PDF), skipping", path.name)
        return RawDocument()

    if figures:
        log.info("%s: %d figure(s) found", path.name, len(figures))
    if bad_pages:
        log.warning("%s: %d unreadable page(s) skipped", path.name, bad_pages)
    return doc


# ---------- .docx ----------

def _docx_rows(table) -> list[list[str]]:
    rows = []
    for row in table.rows:
        seen: list = []
        cells: list[str] = []
        for cell in row.cells:
            duplicate = any(cell._tc is s for s in seen)
            cells.append("" if duplicate else cell.text)
            seen.append(cell._tc)
        rows.append(cells)
    return rows


def _docx_figures(d) -> list[FigureRef]:
    """Embedded images live in the document's media parts, unordered relative to text."""
    figures: list[FigureRef] = []
    try:
        for rel in d.part.rels.values():
            if "image" not in rel.reltype:
                continue
            try:
                blob = rel.target_part.blob
                width, height = 0, 0
                try:                                        # dimensions needed for the size filter
                    from PIL import Image
                    with Image.open(io.BytesIO(blob)) as im:
                        width, height = im.size
                except Exception:
                    pass                                     # unreadable image bytes; keep it, size filter will just let it through
                figures.append(FigureRef(image_bytes=blob, width=width, height=height))
            except Exception:
                continue
    except Exception as e:
        log.warning("figure extraction from .docx failed (%s), continuing without figures", e)
    return figures


def _read_docx(path: Path) -> RawDocument:
    try:
        import docx
        from docx.table import Table as DocxTable
        from docx.text.paragraph import Paragraph
        d = docx.Document(str(path))
    except Exception as e:
        raise InvalidFile(f"{path.name}: cannot open Word file ({e})") from e

    paragraphs: list[str] = []
    last_style = ""
    current_section: str | None = None
    tables: list[Table] = []

    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            text = p.text.strip()
            if not text:
                continue
            try:
                last_style = (p.style.name or "").lower()
            except Exception:
                last_style = ""
            if last_style.startswith("heading") or last_style == "title":
                text = "# " + text
                current_section = p.text.strip()
            paragraphs.append(text)
        elif tag == "tbl":
            caption = None
            if paragraphs:
                caption = is_caption(paragraphs[-1]) or (paragraphs[-1] if last_style == "caption" else None)
                if caption:
                    paragraphs.pop()
            tables.append(Table(rows=_docx_rows(DocxTable(child, d)), caption=caption,
                                section=current_section))

    figures = _docx_figures(d)
    if figures:
        log.info("%s: %d figure(s) found", path.name, len(figures))
    return RawDocument(pages=["\n\n".join(paragraphs)], tables=tables, figures=figures)


# ---------- .csv and .xlsx (the whole file is a table) ----------

def _read_csv(path: Path) -> RawDocument:
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    if not raw.strip():
        return RawDocument()
    try:
        dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(raw, newline=""), dialect))
    return RawDocument(tables=[Table(rows=rows, caption=path.stem.replace("_", " "))])


def _read_xlsx(path: Path) -> RawDocument:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise InvalidFile(f"{path.name}: cannot open Excel file ({e})") from e

    tables: list[Table] = []
    try:
        for ws in wb.worksheets:
            rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
            if any(any(c.strip() for c in r) for r in rows):
                tables.append(Table(rows=rows, caption=ws.title))
    finally:
        wb.close()
    return RawDocument(tables=tables)


# ---------- router ----------

_READERS = {
    ".txt": _read_txt,
    ".md": _read_markdown,
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".csv": _read_csv,
    ".xlsx": _read_xlsx,
}


def read_document(file_path: str | Path) -> RawDocument:
    path = Path(file_path)
    reader = _READERS.get(path.suffix.lower())
    if reader is None:
        raise ParserError(f"No reader for '{path.suffix}'")
    return reader(path)
