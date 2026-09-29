import csv, sys
sys.path.insert(0, '.')
from docx import Document
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.pdfgen import canvas

OUT = "tests/fixtures"
ROWS = [["Model", "Company", "Year"], ["GPT-4", "OpenAI", "2023"], ["Claude 2", "Anthropic", "2023"]]
CAPTION = "Table 1: Model release dates."

with open(f"{OUT}/models.csv", "w", newline="") as f:
    csv.writer(f).writerows(ROWS)

wb = Workbook()
ws = wb.active
ws.title = "Models"
for r in ROWS:
    ws.append(r)
wb.save(f"{OUT}/models.xlsx")

d = Document()
d.add_heading("Company History", level=1)
d.add_paragraph("OpenAI was founded in 2015 in San Francisco and released GPT-4 in March 2023.")
d.add_paragraph(CAPTION)
t = d.add_table(rows=len(ROWS), cols=3)
for i, row in enumerate(ROWS):
    for j, value in enumerate(row):
        t.cell(i, j).text = value
d.save(f"{OUT}/report_with_table.docx")

with open(f"{OUT}/with_table.md", "w") as f:
    f.write(
        "# Company History\n"
        "OpenAI was founded in 2015 in San Francisco.\n\n"
        f"{CAPTION}\n\n"
        "| Model | Company | Year |\n"
        "|---|---|---|\n"
        "| GPT-4 | OpenAI | 2023 |\n"
        "| Claude 2 | Anthropic | 2023 |\n"
    )

styles = getSampleStyleSheet()
grid = Table(ROWS)
grid.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
SimpleDocTemplate(f"{OUT}/table.pdf", pagesize=A4).build([
    Paragraph("OpenAI released GPT-4 in March 2023 and it was widely covered.", styles["Normal"]),
    Spacer(1, 12),
    Paragraph(CAPTION, styles["Normal"]),
    Spacer(1, 6),
    grid,
])

with open(f"{OUT}/fake_docx.docx", "w") as f:
    f.write("just plain text, not a Word file")

# numeric-heavy table
NUM_ROWS = [["Year", "Revenue", "Growth"]] + [[str(2015+i), str(1000*i), f"{i}%"] for i in range(6)]
with open(f"{OUT}/numeric_table.csv", "w", newline="") as f:
    csv.writer(f).writerows(NUM_ROWS)

# a table spanning 2 PDF pages (continuation, no header on page 2)
HEADER = ["Model", "Score"]
PAGE1_ROWS = [HEADER] + [[f"M{i}", str(i)] for i in range(5)]
PAGE2_ROWS = [[f"M{i}", str(i)] for i in range(5, 10)]  # no header row
c = canvas.Canvas(f"{OUT}/spanning_table.pdf", pagesize=A4)
doc = SimpleDocTemplate(f"{OUT}/spanning_table.pdf", pagesize=A4)
from reportlab.platypus import PageBreak
t1 = Table(PAGE1_ROWS); t1.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.5,colors.black)]))
t2 = Table(PAGE2_ROWS); t2.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.5,colors.black)]))
doc.build([Paragraph("Scores table, continues on next page.", styles["Normal"]), t1, PageBreak(), t2])

# a fixture with a heading before a table, for section-inheritance test
d2 = Document()
d2.add_heading("Pricing", level=1)
d2.add_paragraph("Below is the pricing table.")
t2d = d2.add_table(rows=2, cols=2)
t2d.cell(0,0).text = "Plan"; t2d.cell(0,1).text = "Price"
t2d.cell(1,0).text = "Basic"; t2d.cell(1,1).text = "10"
d2.save(f"{OUT}/pricing.docx")

print("done")
