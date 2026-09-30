"""Generates Tier 3 fixtures (figures + a scanned PDF for OCR).
Run from the project root: python make_figure_fixtures.py
Requires: Pillow, PyMuPDF, python-docx (already in requirements.txt from Tier 1/2)
"""
from PIL import Image, ImageDraw, ImageFont
import pymupdf
from docx import Document

OUT = "tests/fixtures"

# a real-sized test "figure"
img = Image.new("RGB", (300, 200), "white")
d = ImageDraw.Draw(img)
d.rectangle([20, 20, 280, 180], outline="black")
d.text((100, 90), "chart", fill="black")
img.save(f"{OUT}/_test_figure.png")

# docx with an embedded figure
doc = Document()
doc.add_heading("Overview", level=1)
doc.add_paragraph("This report includes a chart showing OpenAI's growth.")
doc.add_picture(f"{OUT}/_test_figure.png")
doc.add_paragraph("More text after the figure.")
doc.save(f"{OUT}/report_with_figure.docx")

# pdf with an embedded figure
pdoc = pymupdf.open()
page = pdoc.new_page()
page.insert_text((72, 72), "OpenAI released GPT-4 in March 2023. See the chart below.")
page.insert_image(pymupdf.Rect(72, 100, 372, 300), filename=f"{OUT}/_test_figure.png")
pdoc.save(f"{OUT}/pdf_with_figure.pdf")
pdoc.close()

# a scanned (image-only, no real text layer) PDF, for OCR testing
scan_img = Image.new("RGB", (1000, 300), "white")
sd = ImageDraw.Draw(scan_img)
try:
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 28)
except Exception:
    font = ImageFont.load_default()
sd.text((30, 30), "OpenAI released GPT-4 in March 2023.", fill="black", font=font)
sd.text((30, 80), "The company was co-founded by Sam Altman.", fill="black", font=font)
scan_img.save("/tmp/_scanned_page.png")

sdoc = pymupdf.open()
spage = sdoc.new_page(width=1000, height=300)
spage.insert_image(pymupdf.Rect(0, 0, 1000, 300), filename="/tmp/_scanned_page.png")
sdoc.save(f"{OUT}/scanned_page.pdf")
sdoc.close()

print("Tier 3 fixtures created in", OUT)
